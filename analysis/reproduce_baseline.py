"""Reproduce the documented traditional-BI analyses from the canonical CSV.

This script implements the pre-registered scope in
analysis/BASELINE_REPRODUCTION_PLAN.md. It reads, but never modifies, the
canonical CSV and baseline workbook.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import scipy.stats as stats
import statsmodels.api as sm
from openpyxl import load_workbook


SOURCE = Path("data/Crop_Yield_original.csv")
BASELINE = Path("baseline/Crop Yield Dataset.xlsx")
RESULTS = Path("results")
DETAILS = RESULTS / "baseline_details"
EXPECTED_SHA256 = "58d1866b7a01519593b68c9c929007d10051d070aa957773589f136065cc5f1c"
TOLERANCE = 1e-9
YIELD = "Crop_Yield_ton_per_hectare"
CROPS = ["Cotton", "Maize", "Potato", "Rice", "Sugarcane", "Wheat"]
NUMERIC_COLUMNS = [
    "N", "P", "K", "Soil_pH", "Soil_Moisture", "Organic_Carbon",
    "Temperature", "Humidity", "Rainfall", "Sunlight_Hours", "Wind_Speed",
    "Altitude", "Fertilizer_Used", "Pesticide_Used", YIELD,
]
REGRESSION_PREDICTORS = [
    "N", "P", "K", "Soil_pH", "Soil_Moisture", "Organic_Carbon",
    "Temperature", "Rainfall", "Fertilizer_Used", "Humidity",
    "Sunlight_Hours", "Wind_Speed", "Altitude", "Pesticide_Used",
]


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def json_text(value: Any) -> str:
    def clean(item: Any) -> Any:
        if isinstance(item, dict):
            return {str(k): clean(v) for k, v in item.items()}
        if isinstance(item, (list, tuple)):
            return [clean(v) for v in item]
        if isinstance(item, (np.integer,)):
            return int(item)
        if isinstance(item, (np.floating, float)):
            return None if not np.isfinite(item) else float(item)
        return item

    return json.dumps(clean(value), sort_keys=True, separators=(",", ":"))


def numeric_differences(original: dict[str, Any], reproduced: dict[str, Any]) -> dict[str, float]:
    differences: dict[str, float] = {}

    def walk(left: Any, right: Any, prefix: str = "") -> None:
        if isinstance(left, dict) and isinstance(right, dict):
            for key in left.keys() & right.keys():
                walk(left[key], right[key], f"{prefix}.{key}" if prefix else str(key))
        elif isinstance(left, (int, float)) and isinstance(right, (int, float)):
            differences[prefix] = float(right) - float(left)

    walk(original, reproduced)
    return differences


def differences_match(differences: dict[str, float]) -> bool:
    return all(abs(value) <= TOLERANCE for value in differences.values())


def descriptive_statistics(frame: pd.DataFrame) -> dict[str, dict[str, float]]:
    output: dict[str, dict[str, float]] = {}
    for column in NUMERIC_COLUMNS:
        values = frame[column].astype(float)
        modes = values.mode()
        output[column] = {
            "mean": values.mean(),
            "standard_error": values.sem(ddof=1),
            "median": values.median(),
            "mode": modes.iloc[0],
            "standard_deviation": values.std(ddof=1),
            "sample_variance": values.var(ddof=1),
            "kurtosis": values.kurt(),
            "skewness": values.skew(),
            "range": values.max() - values.min(),
            "minimum": values.min(),
            "maximum": values.max(),
            "sum": values.sum(),
            "count": int(values.count()),
        }
    return output


def extract_original_descriptive(workbook) -> dict[str, dict[str, float]]:
    sheet = workbook["EDA"]
    metric_rows_1 = {
        "mean": 4, "standard_error": 5, "median": 6, "mode": 7,
        "standard_deviation": 8, "sample_variance": 9, "kurtosis": 10,
        "skewness": 11, "range": 12, "minimum": 13, "maximum": 14,
        "sum": 15, "count": 16,
    }
    metric_rows_2 = {
        "mean": 22, "standard_error": 23, "median": 24, "mode": 25,
        "standard_deviation": 26, "sample_variance": 27, "kurtosis": 28,
        "skewness": 29, "range": 30, "minimum": 31, "maximum": 32,
        "sum": 33, "count": 34,
    }
    output: dict[str, dict[str, float]] = {}
    for column_index in range(2, 10):
        name = sheet.cell(2, column_index).value
        output[name] = {metric: sheet.cell(row, column_index).value for metric, row in metric_rows_1.items()}
    for column_index in range(2, 9):
        name = sheet.cell(20, column_index).value
        output[name] = {metric: sheet.cell(row, column_index).value for metric, row in metric_rows_2.items()}
    return output


def crop_summary(frame: pd.DataFrame) -> dict[str, dict[str, float]]:
    grouped = frame.groupby("Crop_Type", sort=True)[YIELD]
    table = grouped.agg(count="count", mean="mean", median="median", standard_deviation="std", minimum="min", maximum="max")
    return {crop: row.to_dict() for crop, row in table.iterrows()}


def original_crop_means(workbook) -> dict[str, dict[str, float]]:
    sheet = workbook["Crop vs Soil vs Irrigation"]
    return {sheet.cell(row, 1).value: {"mean": sheet.cell(row, 6).value} for row in range(3, 9)}


def pearson(frame: pd.DataFrame, x: str, y: str = YIELD) -> dict[str, float]:
    result = stats.pearsonr(frame[x], frame[y])
    return {"r": float(result.statistic), "p_value": float(result.pvalue), "n": int(len(frame))}


def one_way_anova(frame: pd.DataFrame, category: str) -> dict[str, Any]:
    groups = {name: values[YIELD].to_numpy() for name, values in frame.groupby(category, sort=True)}
    result = stats.f_oneway(*groups.values())
    return {
        "f_statistic": float(result.statistic),
        "p_value": float(result.pvalue),
        "n": int(sum(len(values) for values in groups.values())),
        "group_sizes": {name: int(len(values)) for name, values in groups.items()},
    }


def regression(frame: pd.DataFrame) -> dict[str, Any]:
    x = sm.add_constant(frame[REGRESSION_PREDICTORS].astype(float), has_constant="add")
    model = sm.OLS(frame[YIELD].astype(float), x).fit()
    coefficients = {}
    intervals = model.conf_int(alpha=0.05)
    for variable in model.params.index:
        coefficients["Intercept" if variable == "const" else variable] = {
            "coefficient": model.params[variable],
            "standard_error": model.bse[variable],
            "t_statistic": model.tvalues[variable],
            "p_value": model.pvalues[variable],
            "ci_95_lower": intervals.loc[variable, 0],
            "ci_95_upper": intervals.loc[variable, 1],
        }
    return {
        "n": int(model.nobs), "r_squared": model.rsquared,
        "adjusted_r_squared": model.rsquared_adj, "f_statistic": model.fvalue,
        "model_p_value": model.f_pvalue, "coefficients": coefficients,
    }


def extract_original_regression(workbook, crop: str) -> dict[str, Any]:
    sheet = workbook[f"Regression {crop}"]
    locations = {}
    coefficient_header = None
    for row in range(1, 80):
        label = sheet.cell(row, 1).value
        if label in {"R Square", "Adjusted R Square", "Observations"}:
            locations[label] = row
        if sheet.cell(row, 2).value == "Coefficients":
            coefficient_header = row
    if coefficient_header is None:
        raise ValueError(f"Could not locate regression coefficients for {crop}")
    anova_regression_row = next(
        row for row in range(1, 80)
        if sheet.cell(row, 1).value == "Regression" and sheet.cell(row, 5).value is not None
    )
    coefficients = {}
    row = coefficient_header + 1
    while sheet.cell(row, 1).value is not None:
        variable = sheet.cell(row, 1).value
        coefficients[variable] = {
            "coefficient": sheet.cell(row, 2).value,
            "standard_error": sheet.cell(row, 3).value,
            "t_statistic": sheet.cell(row, 4).value,
            "p_value": sheet.cell(row, 5).value,
            "ci_95_lower": sheet.cell(row, 6).value,
            "ci_95_upper": sheet.cell(row, 7).value,
        }
        row += 1
    return {
        "n": int(sheet.cell(locations["Observations"], 2).value),
        "r_squared": sheet.cell(locations["R Square"], 2).value,
        "adjusted_r_squared": sheet.cell(locations["Adjusted R Square"], 2).value,
        "f_statistic": sheet.cell(anova_regression_row, 5).value,
        "model_p_value": sheet.cell(anova_regression_row, 6).value,
        "coefficients": coefficients,
    }


def baseline_regression_input_rows(workbook, crop: str) -> set[tuple[Any, ...]]:
    """Return the copied input rows used on a baseline regression sheet."""
    expected_header = REGRESSION_PREDICTORS + [YIELD]
    rows = list(workbook[f"Regression {crop}"].iter_rows(values_only=True))
    header_index = next(
        index for index, row in enumerate(rows)
        if list(row[: len(expected_header)]) == expected_header
    )
    return {
        tuple(row[: len(expected_header)])
        for row in rows[header_index + 1 :]
        if all(value is not None for value in row[: len(expected_header)])
    }


def write_detail(name: str, rows: list[dict[str, Any]]) -> None:
    pd.DataFrame(rows).to_csv(DETAILS / name, index=False)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--baseline", type=Path, default=BASELINE)
    parser.add_argument("--results", type=Path, default=RESULTS)
    args = parser.parse_args()

    if file_hash(args.source) != EXPECTED_SHA256:
        raise RuntimeError("Canonical dataset hash does not match the approved Phase 1 checksum")
    frame = pd.read_csv(args.source)
    workbook = load_workbook(args.baseline, data_only=True, read_only=True)
    global DETAILS
    DETAILS = args.results / "baseline_details"
    DETAILS.mkdir(parents=True, exist_ok=True)
    generated_at = datetime.now(timezone.utc).isoformat()
    comparisons: list[dict[str, Any]] = []
    register: list[dict[str, Any]] = []

    def record(analysis_id: str, purpose: str, variables: str, method: str,
               filter_text: str, n: int | None, original: Any, reproduced: Any,
               difference: Any, status: str, explanation: str,
               evidence: str, detail_file: str, interpretation: str) -> None:
        comparisons.append({
            "analysis_id": analysis_id, "original_result": json_text(original),
            "reproduced_python_result": json_text(reproduced), "difference": json_text(difference),
            "method": method, "sample_size": n, "status": status,
            "explanation": explanation, "baseline_evidence": evidence,
            "generated_at_utc": generated_at,
        })
        register.append({
            "analysis_id": analysis_id, "business_question": "Baseline reproduction",
            "analysis_purpose": purpose, "dataset": args.source.as_posix(),
            "dataset_sha256": EXPECTED_SHA256, "filters": filter_text,
            "variables": variables, "method": method, "assumptions": "See pre-registered plan",
            "n": n, "code_location": "analysis/reproduce_baseline.py",
            "output_location": f"results/baseline_details/{detail_file}",
            "statistical_result": json_text(reproduced), "interpretation": interpretation,
            "ai_generated": "Code and candidate wording AI-assisted; numbers deterministic Python",
            "independent_ai_review": "NOT_REVIEWED", "human_verification": "PENDING",
            "final_status": "UNVERIFIED",
        })

    # A001
    original_desc = extract_original_descriptive(workbook)
    reproduced_desc = descriptive_statistics(frame)
    desc_diff = numeric_differences(original_desc, reproduced_desc)
    detail_rows = [{"variable": v, "metric": m, "original": original_desc[v][m],
                    "reproduced": reproduced_desc[v][m], "difference": reproduced_desc[v][m] - original_desc[v][m]}
                   for v in original_desc for m in original_desc[v]]
    write_detail("A001_overall_descriptive.csv", detail_rows)
    record("A001", "Reproduce overall descriptive statistics", ";".join(NUMERIC_COLUMNS),
           "Descriptive statistics using pandas sample conventions", "None", len(frame),
           original_desc, reproduced_desc, desc_diff, "MATCH" if differences_match(desc_diff) else "MATERIAL DIFFERENCE",
           "All non-mode workbook metrics agree within floating-point tolerance. Four fields have tied modes: Excel MODE.SNGL returns the first encountered tied value, while pandas mode returns tied values in sorted order and this output records the first sorted value. The report also contains an earlier conflicting narrative estimate (yield 22.2 and fertiliser 206.68) before later reporting the workbook values rounded to 22.34 and 205.99.", "Workbook EDA!A1:I34; report paragraphs 130 and 137",
           "A001_overall_descriptive.csv", "Descriptive benchmark comparison only; no business conclusion.")

    # A002
    original_means = original_crop_means(workbook)
    reproduced_crops = crop_summary(frame)
    crop_diff = numeric_differences(original_means, reproduced_crops)
    write_detail("A002_crop_yield_summary.csv", [{"crop": c, **values} for c, values in reproduced_crops.items()])
    record("A002", "Reproduce crop-level yield summaries", "Crop_Type;" + YIELD,
           "Grouped count, mean, median, sample SD, minimum, maximum", "None", len(frame),
           original_means, reproduced_crops, crop_diff, "MATCH" if differences_match(crop_diff) else "MATERIAL DIFFERENCE",
           "Exact baseline means are available; other summary metrics are reproduced but lack complete stored benchmarks.",
           "Workbook Crop vs Soil vs Irrigation!A3:F8; report paragraphs 130-145",
           "A002_crop_yield_summary.csv", "Crop summaries are descriptive and unverified.")

    # A003 and A004-A009
    overall_fertilizer = pearson(frame, "Fertilizer_Used")
    write_detail("A003_overall_fertilizer_correlation.csv", [overall_fertilizer])
    record("A003", "Reproduce pooled fertiliser-yield relationship", "Fertilizer_Used;" + YIELD,
           "Two-sided Pearson correlation", "None", len(frame),
           {"reported_direction_strength": "very weak; numeric coefficient not located"}, overall_fertilizer, {},
           "UNABLE TO REPRODUCE", "Python result produced, but no traceable numeric pooled benchmark was found.",
           "Report paragraphs 189 and 205", "A003_overall_fertilizer_correlation.csv",
           "Pooled association only; no causal interpretation.")
    fertilizer_sheet = workbook["Correlation Crop and Fertilizer"]
    original_fertilizer = {fertilizer_sheet.cell(row, 1).value: {
        "r": fertilizer_sheet.cell(row, 3).value,
        "p_value": fertilizer_sheet.cell(row, 6).value,
        "n": int(fertilizer_sheet.cell(row, 2).value),
    } for row in range(15, 21)}
    fertilizer_rows = []
    for offset, crop in enumerate(CROPS, 4):
        subset = frame[frame["Crop_Type"] == crop]
        reproduced = pearson(subset, "Fertilizer_Used")
        original = original_fertilizer[crop]
        difference = numeric_differences(original, reproduced)
        fertilizer_rows.append({"analysis_id": f"A{offset:03d}", "crop": crop, **reproduced})
        record(f"A{offset:03d}", f"Reproduce fertiliser-yield correlation for {crop}",
               "Fertilizer_Used;" + YIELD, "Two-sided Pearson correlation", f"Crop_Type={crop}",
               len(subset), original, reproduced, difference,
               "MATCH" if differences_match(difference) else "MATERIAL DIFFERENCE",
               "Compared with the detailed baseline table; an earlier workbook summary contains inconsistent r values for Potato, Rice, and Wheat.",
               f"Workbook Correlation Crop and Fertilizer!A{11+offset}:H{11+offset}",
               "A004_A009_crop_fertilizer_correlations.csv", "Association only; no causal interpretation.")
    write_detail("A004_A009_crop_fertilizer_correlations.csv", fertilizer_rows)

    # A010-A015 irrigation ANOVA
    irrigation_sheet = workbook["ANOVA - Crop type vs Irrigation"]
    irrigation_original = {irrigation_sheet.cell(row, 5).value.replace("Sugacane", "Sugarcane"): irrigation_sheet.cell(row, 6).value for row in range(7, 13)}
    irrigation_rows = []
    for offset, crop in enumerate(CROPS, 10):
        subset = frame[frame["Crop_Type"] == crop]
        reproduced = one_way_anova(subset, "Irrigation_Type")
        original = {"p_value": irrigation_original[crop], "n": len(subset)}
        difference = numeric_differences(original, reproduced)
        irrigation_rows.append({"analysis_id": f"A{offset:03d}", "crop": crop, **{k: json_text(v) if isinstance(v, dict) else v for k, v in reproduced.items()}})
        record(f"A{offset:03d}", f"Reproduce irrigation ANOVA for {crop}", "Irrigation_Type;" + YIELD,
               "One-way ANOVA", f"Crop_Type={crop}", len(subset), original, reproduced, difference,
               "MATCH" if differences_match(difference) else "MATERIAL DIFFERENCE",
               "All four observed irrigation groups retained.", "Workbook ANOVA - Crop type vs Irrigation!E7:G12",
               "A010_A015_irrigation_anova.csv", "Group comparison only; non-significant results retained.")
    write_detail("A010_A015_irrigation_anova.csv", irrigation_rows)

    # A016-A022 rainfall correlations
    overall_rainfall = pearson(frame, "Rainfall")
    original_overall_rainfall = {"r": workbook["Correlation Matrix"].cell(15, 11).value, "n": len(frame)}
    rain_diff = numeric_differences(original_overall_rainfall, overall_rainfall)
    write_detail("A016_overall_rainfall_correlation.csv", [overall_rainfall])
    record("A016", "Reproduce pooled rainfall-yield relationship", "Rainfall;" + YIELD,
           "Two-sided Pearson correlation", "None", len(frame), original_overall_rainfall, overall_rainfall,
           rain_diff, "MATCH" if differences_match(rain_diff) else "MATERIAL DIFFERENCE",
           "Baseline matrix stores r but not p-value.", "Workbook Correlation Matrix!K15",
           "A016_overall_rainfall_correlation.csv", "Association only; no causal interpretation.")
    crop_matrix = workbook["Correlation Matrix by crop"]
    original_rain = {crop_matrix.cell(row, 2).value: crop_matrix.cell(row, 11).value for row in range(2, 8)}
    rain_rows = []
    for offset, crop in enumerate(CROPS, 17):
        subset = frame[frame["Crop_Type"] == crop]
        reproduced = pearson(subset, "Rainfall")
        original = {"r": original_rain[crop], "n": len(subset)}
        difference = numeric_differences(original, reproduced)
        rain_rows.append({"analysis_id": f"A{offset:03d}", "crop": crop, **reproduced})
        record(f"A{offset:03d}", f"Reproduce rainfall-yield correlation for {crop}", "Rainfall;" + YIELD,
               "Two-sided Pearson correlation", f"Crop_Type={crop}", len(subset), original, reproduced,
               difference, "MATCH" if differences_match(difference) else "MATERIAL DIFFERENCE",
               "Baseline matrix stores r but not p-value.", "Workbook Correlation Matrix by crop!B2:Q7",
               "A017_A022_crop_rainfall_correlations.csv", "Association only; no causal interpretation.")
    write_detail("A017_A022_crop_rainfall_correlations.csv", rain_rows)

    # A023-A028 soil ANOVA
    soil_sheet = workbook["ANOVA - Crop type vs Soil Type"]
    soil_original = {soil_sheet.cell(row, 1).value: soil_sheet.cell(row, 2).value for row in range(8, 14)}
    soil_rows = []
    for offset, crop in enumerate(CROPS, 23):
        subset = frame[frame["Crop_Type"] == crop]
        reproduced = one_way_anova(subset, "Soil_Type")
        original = {"p_value": soil_original[crop], "n": len(subset)}
        difference = numeric_differences(original, reproduced)
        soil_rows.append({"analysis_id": f"A{offset:03d}", "crop": crop, **{k: json_text(v) if isinstance(v, dict) else v for k, v in reproduced.items()}})
        record(f"A{offset:03d}", f"Reproduce soil-type ANOVA for {crop}", "Soil_Type;" + YIELD,
               "One-way ANOVA", f"Crop_Type={crop}", len(subset), original, reproduced, difference,
               "MATCH" if differences_match(difference) else "MATERIAL DIFFERENCE",
               "All four observed soil groups retained.", "Workbook ANOVA - Crop type vs Soil Type!A8:C13",
               "A023_A028_soil_anova.csv", "Group comparison only; significant and non-significant results retained.")
    write_detail("A023_A028_soil_anova.csv", soil_rows)

    # A029-A034 regressions
    regression_rows = []
    omitted_regression_rows = []
    for offset, crop in enumerate(CROPS, 29):
        subset = frame[frame["Crop_Type"] == crop]
        reproduced = regression(subset)
        original = extract_original_regression(workbook, crop)
        baseline_inputs = baseline_regression_input_rows(workbook, crop)
        canonical_columns = REGRESSION_PREDICTORS + [YIELD]
        missing_indices = [
            index for index, row in subset[canonical_columns].iterrows()
            if tuple(row.tolist()) not in baseline_inputs
        ]
        for index in missing_indices:
            omitted_regression_rows.append({
                "crop": crop,
                "source_row_id": int(index) + 1,
                "physical_csv_row": int(index) + 2,
                **frame.loc[index, canonical_columns].to_dict(),
            })
        difference = numeric_differences(original, reproduced)
        status = "MATCH" if original["n"] == reproduced["n"] and differences_match(difference) else "MATERIAL DIFFERENCE"
        explanation = ("Full canonical crop subset matches the baseline input population." if original["n"] == reproduced["n"]
                       else f"Baseline regression used {original['n']} rows while the canonical crop subset has {reproduced['n']}. Missing physical CSV rows for this crop: {','.join(str(index + 2) for index in missing_indices)}. The baseline copied input range stops before the final nine source records.")
        regression_rows.append({"analysis_id": f"A{offset:03d}", "crop": crop, "baseline_n": original["n"],
                                "python_n": reproduced["n"], "baseline_r_squared": original["r_squared"],
                                "python_r_squared": reproduced["r_squared"], "r_squared_difference": reproduced["r_squared"] - original["r_squared"],
                                "baseline_adjusted_r_squared": original["adjusted_r_squared"],
                                "python_adjusted_r_squared": reproduced["adjusted_r_squared"], "status": status})
        record(f"A{offset:03d}", f"Reproduce multiple regression for {crop}", ";".join(REGRESSION_PREDICTORS + [YIELD]),
               "OLS with intercept", f"Crop_Type={crop}", len(subset), original, reproduced, difference,
               status, explanation, f"Workbook Regression {crop} summary output",
               "A029_A034_regression_summary.csv", "Conditional associations only; coefficients are not causal effects.")
    write_detail("A029_A034_regression_summary.csv", regression_rows)
    write_detail("A029_A034_regression_omitted_rows.csv", sorted(omitted_regression_rows, key=lambda row: row["physical_csv_row"]))

    # A035 crop by soil means
    original_matrix = {sheet.cell(row, 1).value: {sheet.cell(2, col).value: sheet.cell(row, col).value for col in range(2, 6)}
                       for sheet in [workbook["Crop vs Soil vs Irrigation"]] for row in range(3, 9)}
    reproduced_pivot = frame.pivot_table(index="Crop_Type", columns="Soil_Type", values=YIELD, aggfunc="mean")
    reproduced_matrix = {crop: reproduced_pivot.loc[crop].to_dict() for crop in CROPS}
    matrix_diff = numeric_differences(original_matrix, reproduced_matrix)
    pivot_rows = [{"crop": crop, **reproduced_matrix[crop]} for crop in CROPS]
    write_detail("A035_crop_soil_means.csv", pivot_rows)
    record("A035", "Reproduce crop-by-soil mean yield table", "Crop_Type;Soil_Type;" + YIELD,
           "Arithmetic mean grouped by crop and soil", "None", len(frame), original_matrix, reproduced_matrix,
           matrix_diff, "MATCH" if differences_match(matrix_diff) else "MATERIAL DIFFERENCE",
           "The workbook tab name mentions irrigation, but the stored pivot contains no irrigation dimension.",
           "Workbook Crop vs Soil vs Irrigation!A2:F9", "A035_crop_soil_means.csv",
           "Descriptive group means only; no causal interpretation.")

    # A036-A037 unavailable benchmarks
    record("A036", "Reproduce claimed crop×soil×irrigation comparison", "Crop_Type;Soil_Type;Irrigation_Type;" + YIELD,
           "Not run: baseline numeric output not traceable", "None", None,
           {"reported_result": "qualitative chart/pivot claim"}, {}, {}, "UNABLE TO REPRODUCE",
           "The named workbook output contains crop×soil means only; adding irrigation would create an unsupported benchmark.",
           "Report paragraphs 381 and 412-418; workbook Crop vs Soil vs Irrigation!A1:F9",
           "A036_unable_to_reproduce.csv", "No result generated.")
    record("A037", "Reproduce fertiliser-use-band yield comparison", "Fertilizer_Used;Crop_Type;" + YIELD,
           "Not run: original numeric output not traceable", "None", None,
           {"reported_result": "visual difference across four categories"}, {}, {}, "UNABLE TO REPRODUCE",
           "Traceable band definitions and numeric benchmark were not established from the documented formal output.",
           "Report paragraphs 420-425; workbook EDA table headings",
           "A037_unable_to_reproduce.csv", "No result generated.")
    write_detail("A036_unable_to_reproduce.csv", [{"analysis_id": "A036", "reason": comparisons[-2]["explanation"]}])
    write_detail("A037_unable_to_reproduce.csv", [{"analysis_id": "A037", "reason": comparisons[-1]["explanation"]}])

    pd.DataFrame(comparisons).to_csv(args.results / "baseline_reproduction.csv", index=False, quoting=csv.QUOTE_MINIMAL)
    pd.DataFrame(register).to_csv(args.results / "analysis_register.csv", index=False, quoting=csv.QUOTE_MINIMAL)
    print(json.dumps({"analyses": len(comparisons), "status_counts": pd.Series([r["status"] for r in comparisons]).value_counts().to_dict()}))


if __name__ == "__main__":
    main()

